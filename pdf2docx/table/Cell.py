# -*- coding: utf-8 -*-

'''Table Cell object.
'''


from docx.shared import Pt
from ..common.Element import Element
from ..common.share import rgb_component
from ..common import docx, constants
from ..text.TextBlock import TextBlock
from ..text.Line import Line
from ..text.Lines import Lines
from ..page import Layout # avoid import conflict

class Cell(Element):
    '''Cell object.'''
    def __init__(self, raw:dict=None):
        raw = raw or {}        
        self.bg_color     = raw.get('bg_color', None) # type: int
        self.border_color = raw.get('border_color', (0,0,0,0)) # type: tuple [int]
        self.border_width = raw.get('border_width', (0,0,0,0)) # type: tuple [float]
        self.merged_cells = raw.get('merged_cells', (1,1)) # type: tuple [int]

        # layout
        self.layout = Layout.Layout(parent=self).restore(raw)

        super().__init__(raw)


    @property
    def text(self):
        '''Text contained in this cell.'''
        return '\n'.join([block.text for block in self.layout.blocks]) if bool(self) else None


    @property
    def working_bbox(self):
        '''Bbox with border considered.'''
        x0, y0, x1, y1 = self.bbox
        w_top, w_right, w_bottom, w_left = self.border_width
        return (x0+w_left/2.0, y0+w_top/2.0, x1-w_right/2.0, y1-w_bottom/2.0)

    
    def compare(self, cell, threshold:float=0.9):
        """Whether has same structure with given Cell.

        Args:
            cell (Cell): Target cell to compare.
            threshold (float, optional): Considered as same bbox if exceeds this value. Defaults to 0.9.

        Returns:
            tuple: ``(True or False, message)``
        """
        # bbox
        res, msg = super().compare(cell, threshold)
        if not res: return res, msg

        # cell style        
        if self.bg_color != cell.bg_color:
            return False, f'Inconsistent background color @ Cell {self.bbox}:\n{self.bg_color} v.s. {cell.bg_color} (expected)'

        if tuple(self.border_color) != tuple(cell.border_color):
            return False, f'Inconsistent border color @ Cell {self.bbox}:\n{self.border_color} v.s. {cell.border_color} (expected)'

        if tuple(self.border_width) != tuple(cell.border_width):
            return False, f'Inconsistent border width @ Cell {self.bbox}:\n{self.border_width} v.s. {cell.border_width} (expected)'

        if tuple(self.merged_cells) != tuple(cell.merged_cells):
            return False, f'Inconsistent count of merged cells @ Cell {self.bbox}:\n{self.merged_cells} v.s. {cell.merged_cells} (expected)'

        return True, ''


    def store(self):
        if bool(self):
            res = super().store()
            res.update({
                'bg_color': self.bg_color,
                'border_color': self.border_color,
                'border_width': self.border_width,
                'merged_cells': self.merged_cells,
                'blocks': self.layout.blocks.store()
            })
            return res
        else:
            return None


    def plot(self, page):
        '''Plot cell and its sub-layout.'''        
        super().plot(page)
        self.layout.blocks.plot(page)


    def add_block(self, block):
        '''Add block to this cell. 
        
        Args:
            block (TextBlock, TableBlock): Text/table block to add.
        
        .. note::
            If a text block is partly contained in a cell, it must deep into line -> span -> char.
        '''
        # add block directly if fully contained in cell
        if self.contains(block, constants.FACTOR_ALMOST):
            self.layout.blocks.append(block)
            return
        
        # add nothing if no intersection
        if not self.bbox & block.bbox: return

        # otherwise, further check lines in text block
        if not block.is_text_image_block():  return
        
        # NOTE: add each line as a single text block to avoid overlap between table block and combined lines
        split_block = TextBlock()
        lines = [line.intersects(self.bbox) for line in block.lines]
        split_block.add(lines)
        self.layout.blocks.append(split_block)


    def add_shape(self, shape):
        '''Add shape to this cell. 
        
        Args:
            shape (Shape): Shape to add.
        '''
        # add shape if contained in cell
        if self.bbox & shape.bbox:
            self.layout.shapes.append(shape)


    def set_stream_table_layout(self, settings:dict):
        '''Set stream table layout to ensure any float layout converted to flow layout.'''
        # create nest table if float layout still exists
        from .TablesConstructor import TablesConstructor
        from .TableStructure import TableStructure

        # bbox range of stream table
        inner_bbox, outer_bbox = self.working_bbox, self.working_bbox
        outer_borders = TablesConstructor._outer_borders(inner_bbox, outer_bbox)

        # stream table contents        
        def sub_lines(block): # get sub-lines from block
            return block.lines if block.is_text_image_block() else [Line().update_bbox(block.bbox)]
        table_lines = Lines()
        for block in self.layout.blocks:
            table_lines.extend(sub_lines(block))

        # parse stream borders
        strokes = TablesConstructor.stream_strokes(table_lines, outer_borders, explicit_strokes=[], explicit_shadings=[])
        if not strokes: return

        # parse table structure
        strokes.sort_in_reading_order() # required
        table = TableStructure(strokes, settings).parse(fills=[]).to_table_block()
        if not table: return

        # parse table content
        table.set_stream_table_block()
        self.layout.blocks.assign_table_contents([table], settings)


    def make_docx(self, table, indexes):
        '''Set cell style and assign contents.
        
        Args:
            table (Table): ``python-docx`` table instance.
            indexes (tuple): Row and column indexes, ``(i, j)``.
        '''        
        # set cell style, e.g. border, shading, cell width
        self._set_style(table, indexes)
        
        # ignore merged cells
        if not bool(self):  return

        # merge cells
        n_row, n_col = self.merged_cells
        i, j = indexes
        docx_cell = table.cell(i, j)
        if n_row*n_col!=1:
            _cell = table.cell(i+n_row-1, j+n_col-1)
            docx_cell.merge(_cell)
        
        # ---------------------
        # cell width (cell height is set by row height)
        # ---------------------
        # experience: width of merged cells may change if not setting width for merged cells
        x0, y0, x1, y1 = self.bbox
        docx_cell.width = Pt(x1-x0)

        # insert contents
        # NOTE: there exists an empty paragraph already in each cell, which should be deleted first to
        # avoid unexpected layout. `docx_cell._element.clear_content()` works here.
        # But, docx requires at least one paragraph in each cell, otherwise resulting in a repair error. 
        if self.layout.blocks:
            docx_cell._element.clear_content()
            self.layout.blocks.make_docx(docx_cell)


    def _set_style(self, table, indexes):
        '''Set ``python-docx`` cell style, e.g. border, shading, width, row height, 
        based on cell block parsed from PDF.
        
        Args:
            table (Table): ``python-docx`` table object.
            indexes (tuple): ``(i, j)`` index of current cell in table.
        '''
        i, j = indexes
        docx_cell = table.cell(i, j)
        n_row, n_col = self.merged_cells

        # ---------------------
        # border style
        # ---------------------
        # NOTE: border width is specified in eighths of a point, with a minimum value of 
        # two (1/4 of a point) and a maximum value of 96 (twelve points)
        keys = ('top', 'end', 'bottom', 'start')
        kwargs = {}
        for k, w, c in zip(keys, self.border_width, self.border_color):
            # skip if width=0 -> will not show in docx
            if not w: continue

            hex_c = f'#{hex(c)[2:].zfill(6)}'
            kwargs[k] = {
                'sz': 8*w, 'val': 'single', 'color': hex_c.upper()
            }

        # merged cells are assumed to have same borders with the main cell        
        for m in range(i, i+n_row):
            for n in range(j, j+n_col):
                docx.set_cell_border(table.cell(m, n), **kwargs)        

        # ---------------------
        # cell bg-color
        # ---------------------
        if self.bg_color!=None:
            docx.set_cell_shading(docx_cell, self.bg_color)
        
        # ---------------------
        # clear cell margin
        # ---------------------
        # NOTE: the start position of a table is based on text in cell, rather than left border of table. 
        # They're almost aligned if left-margin of cell is zero.
        docx.set_cell_margins(docx_cell, start=0, end=0)

        # set vertical direction if contained text blocks are in vertical direction
        if self.layout.blocks.is_vertical_text:
            docx.set_vertical_cell_direction(docx_cell)